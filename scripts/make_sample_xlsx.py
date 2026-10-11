"""Generate the Excel sample declaration (synthetic data, fictional supplier).

Run from the repo root: python scripts/make_sample_xlsx.py

openpyxl stamps the current time into the file, so the output is rewritten
with fixed timestamps. Stable bytes keep the extraction cache valid.
"""

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / "samples" / "08_pellham_blank_conc.xlsx"
FIXED = datetime(2026, 7, 1, 9, 0, 0)
STAMP = FIXED.strftime("%Y-%m-%dT%H:%M:%SZ").encode()

ROWS = [
    ("Lead", "7439-92-1", 0.002, "% w/w"),
    ("Mercury", "7439-97-6", 0.0001, "% w/w"),
    ("Cadmium", "7440-43-9", None, "% w/w"),
    ("Hexavalent chromium", "18540-29-9", None, "% w/w"),
    ("PBB", "59536-65-1", 0.0001, "% w/w"),
    ("PBDE", "1163-19-5", 0.0001, "% w/w"),
    ("DEHP", "117-81-7", 0.001, "% w/w"),
    ("BBP", "85-68-7", 0.001, "% w/w"),
    ("DBP", "84-74-2", 0.001, "% w/w"),
    ("DIBP", "84-69-5", 0.001, "% w/w"),
]


def build() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "RoHS Declaration"
    ws["A1"] = "Synthetic sample for the RADAR prototype. Fictional supplier."
    ws["A1"].font = Font(italic=True)
    ws["A2"], ws["B2"] = "Supplier", "Pellham Fasteners (fictional)"
    ws["A3"], ws["B3"] = "Component", "PF-16 zinc-plated steel screw, M4"
    ws["A4"], ws["B4"] = "Date", "2026-07-01"
    header = ("Substance", "CAS No.", "Concentration", "Unit")
    ws.append([])
    ws.append(header)
    for cell in ws[6]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E6E9EE")
    for row in ROWS:
        ws.append(row)
    ws.append([])
    ws.append(["Values per homogeneous material."])
    for col, width in zip("ABCD", (24, 14, 15, 10)):
        ws.column_dimensions[col].width = width
    wb.properties.created = FIXED
    wb.properties.modified = FIXED

    raw = io.BytesIO()
    wb.save(raw)
    # Re-zip with a fixed timestamp per entry so the bytes never change between runs.
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl overwrites "modified" with the save time; pin it.
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + STAMP, data)
            dst.writestr(fixed, data)
    return out.getvalue()


if __name__ == "__main__":
    OUT.write_bytes(build())
    print(f"wrote samples/{OUT.name}")
