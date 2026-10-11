"""Turn a declaration file into plain text for the model to read.

Supported: .md, .txt, .csv, .pdf, .xlsx. No API calls here.

Spreadsheets are rendered one row per line with cells separated by " | ".
Empty cells inside the table are written as "(blank)" so the model can see
that a value is missing rather than guessing it was never there.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_TYPES = (".md", ".txt", ".csv", ".pdf", ".xlsx")
BLANK = "(blank)"


class UnsupportedDocument(ValueError):
    pass


@dataclass(frozen=True)
class Document:
    filename: str
    kind: str      # file extension without the dot
    data: bytes    # original bytes; the cache key is their hash
    text: str      # what the model reads


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(f"[Page {i}]\n{t.strip()}" for i, t in enumerate(pages, start=1))


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # repr keeps the digits exactly as stored (0.0001, not 1e-04 rounding surprises)
        text = repr(value)
        return text[:-2] if text.endswith(".0") else text
    return str(value).strip()


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    blocks = []
    for ws in wb.worksheets:
        rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        # Drop trailing empty columns and fully empty rows; keep row numbers for traceability.
        width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
        lines = [f"[Sheet: {ws.title}]"]
        for n, row in enumerate(rows, start=1):
            cells = (row + [""] * width)[:width]
            filled = [i for i, c in enumerate(cells) if c]
            if not filled:
                continue
            # A table row fills most of the table width: pad it so a blank last cell still shows.
            # Shorter rows (titles, "Supplier | name") only show blanks between filled cells.
            is_table_row = len(filled) * 2 > width
            start, end = (0, width) if is_table_row else (filled[0], filled[-1] + 1)
            lines.append(f"Row {n}: " + " | ".join(c or BLANK for c in cells[start:end]))
        blocks.append("\n".join(lines))
    wb.close()
    return "\n\n".join(blocks)


def load_document(filename: str, data: bytes) -> Document:
    kind = Path(filename).suffix.lower()
    if kind not in SUPPORTED_TYPES:
        raise UnsupportedDocument(f"{filename}: unsupported type '{kind}'. Use {', '.join(SUPPORTED_TYPES)}")
    if kind == ".pdf":
        text = _pdf_text(data)
    elif kind == ".xlsx":
        text = _xlsx_text(data)
    else:
        text = _decode(data)
    if not text.strip():
        raise UnsupportedDocument(f"{filename}: no readable text found")
    return Document(filename=Path(filename).name, kind=kind.lstrip("."), data=data, text=text.strip())


def load_path(path: str | Path) -> Document:
    path = Path(path)
    return load_document(path.name, path.read_bytes())
